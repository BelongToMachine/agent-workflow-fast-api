import io
import json
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from app.services.document_parsing import (
    ParsedDocument,
    paginate_parsed_document,
    parse_document,
)


def _assert_contract(document: ParsedDocument) -> None:
    assert document.schema_version == "parsed-document.v1"
    assert document.file_hash
    assert document.original_name
    assert document.mime_type
    assert document.parser
    assert document.parser_version
    assert document.content_type in {"text", "structured", "mixed"}
    assert document.blocks
    assert all(block.block_id for block in document.blocks)
    assert all(block.kind for block in document.blocks)
    assert all(isinstance(block.text, str) for block in document.blocks)
    assert all(block.locator for block in document.blocks)


def test_text_and_markdown_documents_have_stable_blocks() -> None:
    text = parse_document("notes.txt", b"First line\r\n\r\nSecond line")
    markdown = parse_document(
        "notes.md",
        b"# Product\n\nA short description.\n\n| Name | Price |\n| --- | ---: |\n| Chair | 10 |",
    )

    _assert_contract(text)
    _assert_contract(markdown)
    assert text.blocks[0].locator == {"lineStart": 1, "lineEnd": 3}
    assert markdown.blocks[0].kind == "heading"
    assert markdown.blocks[0].locator == {"lineStart": 1, "lineEnd": 1}
    assert any(block.kind == "table" for block in markdown.blocks)

    repeat = parse_document("notes.txt", b"First line\r\n\r\nSecond line")
    assert [block.block_id for block in text.blocks] == [
        block.block_id for block in repeat.blocks
    ]


def test_json_documents_preserve_json_paths_and_canonical_text() -> None:
    document = parse_document(
        "products.json",
        json.dumps([{"name": "Chair", "price": 10}, {"name": "Desk"}]).encode(),
    )

    _assert_contract(document)
    assert document.content_type == "structured"
    assert document.blocks[0].locator == {"jsonPath": "$[0]"}
    assert document.blocks[0].data == {"name": "Chair", "price": 10}
    assert '"name": "Chair"' in document.blocks[0].text


def test_csv_documents_emit_header_and_row_records() -> None:
    document = parse_document("products.csv", b"name,price\nChair,10\nDesk,20\n")

    _assert_contract(document)
    assert document.content_type == "structured"
    assert document.blocks[0].kind == "table_header"
    assert document.blocks[0].locator == {"row": 1}
    assert document.blocks[1].kind == "table_row"
    assert document.blocks[1].locator == {"row": 2}
    assert document.blocks[1].data == {"name": "Chair", "price": "10"}


def test_xlsx_documents_emit_sheet_and_row_records() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Products"
    worksheet.append(["name", "price"])
    worksheet.append(["Chair", 10])
    stream = io.BytesIO()
    workbook.save(stream)

    document = parse_document("products.xlsx", stream.getvalue())

    _assert_contract(document)
    assert document.content_type == "structured"
    assert document.blocks[0].kind == "sheet"
    assert document.blocks[0].text == "[Sheet: Products]"
    assert document.blocks[1].kind == "table_header"
    assert document.blocks[1].locator == {"sheet": "Products", "row": 1}
    assert document.blocks[2].data == {"name": "Chair", "price": 10}


def test_pdf_documents_emit_page_and_paragraph_locators(monkeypatch) -> None:
    class FakePage:
        def __init__(self, content: str) -> None:
            self.content = content

        def extract_text(self) -> str:
            return self.content

    monkeypatch.setattr(
        "app.services.document_parsing.PdfReader",
        lambda _stream: SimpleNamespace(
            pages=[FakePage("Product brief\n\nChair details"), FakePage("")]
        ),
    )

    document = parse_document("brief.pdf", b"%PDF-fake")

    _assert_contract(document)
    assert document.content_type == "text"
    assert document.blocks[0].locator == {"page": 1, "paragraph": 1}
    assert any("no extractable text" in warning for warning in document.warnings)


def test_pptx_documents_emit_slide_shape_and_table_locators() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text = "Riboton product brief"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(2)).table
    table.cell(0, 0).text = "Product"
    table.cell(0, 1).text = "Price"
    table.cell(1, 0).text = "Riboton"
    table.cell(1, 1).text = "99"
    stream = io.BytesIO()
    presentation.save(stream)

    document = parse_document("brief.pptx", stream.getvalue())

    _assert_contract(document)
    assert document.content_type == "mixed"
    assert any(block.locator == {"slide": 1, "shape": 1} for block in document.blocks)
    assert any(
        block.locator == {"slide": 1, "shape": 2, "row": 2}
        and block.data == {"Product": "Riboton", "Price": "99"}
        for block in document.blocks
    )


def test_pagination_can_return_text_without_structured_payload() -> None:
    document = parse_document("products.csv", b"name,price\nChair,10\nDesk,20\n")

    page = paginate_parsed_document(document, output="text", offset=1, limit=1)

    assert page.total_blocks == 3
    assert page.next_cursor == 2
    assert page.truncated is True
    assert len(page.blocks) == 1
    assert page.blocks[0].text == "Chair\t10"
    assert page.blocks[0].data is None


def test_legacy_ppt_extension_is_rejected() -> None:
    with pytest.raises(ValueError, match="Unsupported knowledge file type"):
        parse_document("legacy.ppt", b"presentation")
