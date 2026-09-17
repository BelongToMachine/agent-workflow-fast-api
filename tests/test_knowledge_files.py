import asyncio
import io
import json
from contextlib import asynccontextmanager
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from uuid import UUID

import pytest
from fastapi import UploadFile
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pydantic import ValidationError
from starlette.datastructures import Headers

from app.api.routes.knowledge_files import (
    FILE_INSERT_QUERY,
    FILE_LIST_QUERY,
    PARSED_DOCUMENT_BY_FILE_QUERY,
    KnowledgeChunkEmbeddingRequest,
    KnowledgeChunkListResponse,
    KnowledgeFileSummary,
    KnowledgeParsedDocumentListResponse,
    KnowledgeParsedDocumentResponse,
    _chunk_text,
    _content_matches_extension,
    _extract_text,
    _safe_filename,
    _storage_path,
    embed_knowledge_file_chunks,
    get_parsed_knowledge_document,
    list_knowledge_file_chunks,
    list_parsed_knowledge_documents,
    materialize_knowledge_chunks,
    materialize_knowledge_file_chunks,
    parse_knowledge_file,
    process_knowledge_file,
    upload_knowledge_file,
)
from app.core.auth import AuthenticatedUser
from app.core.config import MAX_KNOWLEDGE_FILE_BYTES, Settings, get_settings
from app.db.errors import DatabaseUnavailableError
from app.db.migrate_knowledge_ingestion import MIGRATION_PATH
from app.main import app
from app.services.document_parsing import parse_document
from app.services.storage import LocalKnowledgeStorage

client = TestClient(app)


class FakeResult:
    def __init__(
        self,
        row: dict[str, object] | None = None,
        *,
        scalar_value: object | None = None,
    ) -> None:
        self.row = row
        self.scalar_value = scalar_value

    def mappings(self):
        return self

    def first(self):
        return self.row

    def one(self):
        assert self.row is not None
        return self.row

    def scalar_one_or_none(self):
        return self.scalar_value


class FakeTransaction:
    async def __aenter__(self):
        return self

    async def __aexit__(self, *_args) -> None:
        return None


class FakeIngestionConnection:
    def __init__(self, row: dict[str, object]) -> None:
        self.row = row
        self.status_updates: list[dict[str, object]] = []
        self.inserted_chunks: list[dict[str, object]] = []
        self.parsed_documents: list[dict[str, object]] = []
        self.chunk_status_updates: list[dict[str, object]] = []

    async def execute(self, statement, parameters=None):
        sql = str(statement)
        if 'SELECT' in sql and 'KnowledgeParsedDocument' in sql:
            return FakeResult(self.parsed_documents[-1] if self.parsed_documents else None)
        if sql.lstrip().startswith("SELECT"):
            return FakeResult(self.row)
        if 'INSERT INTO "KnowledgeParsedDocument"' in sql:
            self.parsed_documents.append(parameters or {})
            return FakeResult()
        if 'UPDATE "KnowledgeParsedDocument"' in sql:
            self.chunk_status_updates.append(parameters or {})
            if 'RETURNING "id"' in sql:
                return FakeResult(scalar_value=parameters.get("parsed_document_id"))
        if 'UPDATE "KnowledgeFile"' in sql:
            self.status_updates.append(parameters or {})
        if 'INSERT INTO "KnowledgeChunk"' in sql:
            self.inserted_chunks.extend(parameters or [])
        return FakeResult()

    def begin(self):
        return FakeTransaction()

    async def rollback(self) -> None:
        return None


class FakeConnectionContext:
    def __init__(self, connection: FakeIngestionConnection) -> None:
        self.connection = connection

    async def __aenter__(self):
        return self.connection

    async def __aexit__(self, *_args) -> None:
        return None


class FakeUploadResult:
    def __init__(
        self,
        *,
        row: dict[str, object] | None = None,
        rows: list[dict[str, object]] | None = None,
        scalar_value: object | None = None,
    ) -> None:
        self.row = row
        self.rows = rows or ([] if row is None else [row])
        self.scalar_value = scalar_value

    def mappings(self):
        return self

    def first(self):
        return self.row

    def all(self):
        return self.rows

    def one(self) -> dict[str, object]:
        assert self.row is not None
        return self.row

    def scalar_one_or_none(self) -> object | None:
        return self.scalar_value


class FakeUploadConnection:
    def __init__(self, *, inserted: bool = True) -> None:
        self.inserted = inserted
        self.calls: list[tuple[str, dict[str, object]]] = []
        self.file_row: dict[str, object] = {
            "byte_size": 0,
            "created_at": datetime(2026, 8, 17, 12, 30),
            "error_message": None,
            "file_hash": "",
            "file_id": None,
            "knowledge_base_id": UUID("00000000-0000-0000-0000-000000000002"),
            "mime_type": "text/csv",
            "original_name": "data.csv",
            "status": "pending",
            "storage_provider": "local",
            "updated_at": datetime(2026, 8, 17, 12, 30),
            "workspace_id": UUID("00000000-0000-0000-0000-000000000001"),
        }

    async def execute(self, query: object, params: dict[str, object]) -> FakeUploadResult:
        sql = str(query)
        self.calls.append((sql, params))
        if "INSERT INTO \"KnowledgeFile\"" in sql:
            if self.inserted:
                self.file_row.update(
                    {
                        "byte_size": params["byte_size"],
                        "file_hash": params["file_hash"],
                        "file_id": params["file_id"],
                        "original_name": params["original_name"],
                        "storage_provider": params["storage_provider"],
                        "workspace_id": params["workspace_id"],
                    }
                )
                return FakeUploadResult(scalar_value=params["file_id"])
            return FakeUploadResult(scalar_value=None)
        if "SELECT" in sql:
            return FakeUploadResult(row=self.file_row)
        return FakeUploadResult()

    def begin(self) -> FakeTransaction:
        return FakeTransaction()

    async def rollback(self) -> None:
        return None


def upload_connection_context(connection: FakeUploadConnection):
    @asynccontextmanager
    async def context():
        yield connection

    return context


class FakeUploadStorage:
    provider = "local"

    def __init__(self) -> None:
        self.uploads: list[tuple[str, bytes]] = []
        self.deleted: list[str] = []

    async def put(self, storage_key: str, content: bytes) -> None:
        self.uploads.append((storage_key, content))

    async def delete(self, storage_key: str) -> None:
        self.deleted.append(storage_key)


class FakeBackgroundTasks:
    def __init__(self) -> None:
        self.tasks: list[tuple[object, tuple[object, ...]]] = []

    def add_task(self, function, *args, **_kwargs) -> None:
        self.tasks.append((function, args))


@pytest.fixture
def ingestion_disabled_settings():
    settings = Settings(
        environment="development",
        auth_secret="code-secret",
        knowledge_ingestion_enabled=False,
    )
    app.dependency_overrides[get_settings] = lambda: settings
    yield settings
    app.dependency_overrides.pop(get_settings, None)


def test_knowledge_file_summary_requires_storage_metadata() -> None:
    with pytest.raises(ValidationError):
        KnowledgeFileSummary.model_validate(
            {
                "byteSize": 10,
                "createdAt": "2026-08-17T00:00:00Z",
                "fileHash": "abc",
                "fileId": str(UUID(int=1)),
                "knowledgeBaseId": str(UUID(int=2)),
                "mimeType": "text/csv",
                "originalName": "data.csv",
                "status": "pending",
                "updatedAt": "2026-08-17T00:00:00Z",
                "workspaceId": str(UUID(int=3)),
            }
        )


def test_file_parser_supports_csv_and_xlsx() -> None:
    assert "name\tprice" in _extract_text("data.csv", b"name,price\nchair,10")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(["name", "price"])
    sheet.append(["chair", 10])
    output = io.BytesIO()
    workbook.save(output)

    extracted = _extract_text("data.xlsx", output.getvalue())
    assert "[Sheet: Products]" in extracted
    assert "chair\t10" in extracted


def test_file_parser_supports_pptx_text_and_tables() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text = "Riboton product brief"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Product"
    table.cell(0, 1).text = "Price"
    table.cell(1, 0).text = "Riboton"
    table.cell(1, 1).text = "99"

    output = io.BytesIO()
    presentation.save(output)

    extracted = _extract_text("2026 Riboton.pptx", output.getvalue())
    assert "[Slide: 1]" in extracted
    assert "Riboton product brief" in extracted
    assert "Product\tPrice" in extracted
    assert "Riboton\t99" in extracted


def test_binary_file_signatures_match_the_declared_extension() -> None:
    assert _content_matches_extension(".pdf", b"%PDF-1.7 content")
    assert _content_matches_extension(".xlsx", b"PK\x03\x04workbook")
    assert _content_matches_extension(".pptx", b"PK\x03\x04presentation")
    assert not _content_matches_extension(".pdf", b"not-a-pdf")
    assert not _content_matches_extension(".xlsx", b"not-a-workbook")
    assert _content_matches_extension(".csv", b"name,price\nchair,10")


def test_knowledge_file_limit_is_100_mib_and_cannot_be_raised_further() -> None:
    assert Settings(
        knowledge_max_file_bytes=MAX_KNOWLEDGE_FILE_BYTES
    ).knowledge_max_file_bytes == 100 * 1024 * 1024
    with pytest.raises(ValidationError):
        Settings(knowledge_max_file_bytes=MAX_KNOWLEDGE_FILE_BYTES + 1)


def test_file_name_and_storage_path_are_sandboxed(tmp_path) -> None:
    assert _safe_filename("../../secret file.csv") == "secret_file.csv"
    settings = Settings(knowledge_storage_dir=str(tmp_path))

    with pytest.raises(Exception):
        _storage_path(settings, "../../secret.txt")


def test_knowledge_filename_preserves_language_and_extension() -> None:
    assert _safe_filename("产品调研.xlsx") == "产品调研.xlsx"
    long_name = _safe_filename("产品" * 200 + ".xlsx")
    assert long_name.endswith(".xlsx")
    assert len(long_name.encode("utf-8")) <= 160


def test_chunking_adds_overlap_without_empty_chunks() -> None:
    chunks = _chunk_text("a" * 2500)

    assert len(chunks) == 3
    assert all(chunks)
    assert chunks[0][-120:] == chunks[1][:120]


def test_process_knowledge_file_persists_parsed_document_without_chunks(
    monkeypatch,
    tmp_path,
) -> None:
    file_id = UUID("00000000-0000-0000-0000-000000000010")
    workspace_id = UUID("00000000-0000-0000-0000-000000000001")
    knowledge_base_id = UUID("00000000-0000-0000-0000-000000000002")
    storage_key = "workspace/knowledge/data.csv"
    row = {
        "storage_provider": "local",
        "storage_key": storage_key,
        "original_name": "data.csv",
        "knowledge_base_id": knowledge_base_id,
    }
    settings = Settings(
        knowledge_ingestion_enabled=True,
        knowledge_embeddings_enabled=False,
        knowledge_storage_dir=str(tmp_path),
    )
    storage = LocalKnowledgeStorage(str(tmp_path))
    connection = FakeIngestionConnection(row)

    asyncio.run(storage.put(storage_key, b"name,price\nchair,10"))
    monkeypatch.setattr("app.api.routes.knowledge_files.get_settings", lambda: settings)
    monkeypatch.setattr(
        "app.api.routes.knowledge_files.get_knowledge_storage_for_provider",
        lambda _settings, _provider: storage,
    )
    monkeypatch.setattr(
        "app.api.routes.knowledge_files.get_db_connection",
        lambda: FakeConnectionContext(connection),
    )

    asyncio.run(process_knowledge_file(file_id, workspace_id))

    assert [update["status"] for update in connection.status_updates] == [
        "processing",
        "ready",
    ]
    assert connection.inserted_chunks == []
    assert len(connection.parsed_documents) == 1
    assert connection.parsed_documents[0]["file_id"] == file_id
    assert len(connection.parsed_documents[0]["file_hash"]) == 64
    document = json.loads(connection.parsed_documents[0]["document"])
    assert document["schemaVersion"] == "parsed-document.v1"
    parsed_text = "\n".join(block["text"] for block in document["blocks"])
    assert "name\tprice" in parsed_text
    assert "chair\t10" in parsed_text


def test_materialize_knowledge_chunks_is_a_separate_explicit_step(
    monkeypatch,
    tmp_path,
) -> None:
    file_id = UUID("00000000-0000-0000-0000-000000000010")
    workspace_id = UUID("00000000-0000-0000-0000-000000000001")
    knowledge_base_id = UUID("00000000-0000-0000-0000-000000000002")
    parsed_document_id = UUID("00000000-0000-0000-0000-000000000011")
    storage_key = "workspace/knowledge/data.csv"
    storage = LocalKnowledgeStorage(str(tmp_path))
    parsed_document = parse_document(
        "data.csv",
        b"name,price\nchair,10",
        file_id=str(file_id),
        file_hash="hash-123",
        mime_type="text/csv",
    )
    row = {
        "storage_provider": "local",
        "storage_key": storage_key,
        "original_name": "data.csv",
        "knowledge_base_id": knowledge_base_id,
        "file_hash": "hash-123",
    }
    connection = FakeIngestionConnection(row)
    connection.parsed_documents.append(
        {
            "parsed_document_id": parsed_document_id,
            "file_id": file_id,
            "knowledge_base_id": knowledge_base_id,
            "workspace_id": workspace_id,
            "document": parsed_document.model_dump(by_alias=True),
            "file_hash": "hash-123",
            "parser": parsed_document.parser,
            "parser_version": parsed_document.parser_version,
            "chunk_status": "processing",
            "chunk_error_message": None,
            "created_at": datetime(2026, 8, 17, 12, 30),
            "updated_at": datetime(2026, 8, 17, 12, 30),
        }
    )

    asyncio.run(storage.put(storage_key, b"name,price\nchair,10"))
    settings = Settings(
        knowledge_ingestion_enabled=True,
        knowledge_embeddings_enabled=True,
        knowledge_storage_dir=str(tmp_path),
    )
    embedding_calls: list[list[str]] = []

    async def unexpected_embed(texts, _settings):
        embedding_calls.append(texts)
        return [[0.1] * 1024 for _ in texts]

    monkeypatch.setattr("app.api.routes.knowledge_files.embed_texts", unexpected_embed)
    monkeypatch.setattr("app.api.routes.knowledge_files.get_settings", lambda: settings)
    monkeypatch.setattr(
        "app.api.routes.knowledge_files.get_knowledge_storage_for_provider",
        lambda _settings, _provider: storage,
    )
    monkeypatch.setattr(
        "app.api.routes.knowledge_files.get_db_connection",
        lambda: FakeConnectionContext(connection),
    )

    asyncio.run(
        materialize_knowledge_chunks(file_id, workspace_id, parsed_document_id)
    )

    assert [chunk["content"] for chunk in connection.inserted_chunks] == [
        "name\tprice\nchair\t10"
    ]
    assert embedding_calls == []
    assert "embedding" not in connection.inserted_chunks[0]
    metadata = json.loads(connection.inserted_chunks[0]["metadata"])
    assert metadata["fileHash"] == "hash-123"
    assert metadata["parsedDocumentId"] == str(parsed_document_id)
    assert connection.chunk_status_updates[-1]["chunk_status"] == "ready"


def test_materialize_knowledge_chunks_marks_failed_after_database_connection_error(
    monkeypatch,
) -> None:
    file_id = UUID("00000000-0000-0000-0000-000000000010")
    workspace_id = UUID("00000000-0000-0000-0000-000000000001")
    parsed_document_id = UUID("00000000-0000-0000-0000-000000000011")
    status_connection = FakeIngestionConnection({})

    class UnavailableConnectionContext:
        async def __aenter__(self):
            raise DatabaseUnavailableError()

        async def __aexit__(self, *_args) -> None:
            return None

    connection_contexts = iter(
        [
            UnavailableConnectionContext(),
            UnavailableConnectionContext(),
            FakeConnectionContext(status_connection),
        ]
    )
    monkeypatch.setattr(
        "app.api.routes.knowledge_files.get_db_connection",
        lambda: next(connection_contexts),
    )

    asyncio.run(
        materialize_knowledge_chunks(file_id, workspace_id, parsed_document_id)
    )

    assert status_connection.chunk_status_updates == [
        {
            "chunk_error_message": "The database is temporarily unavailable.",
            "chunk_status": "failed",
            "file_id": file_id,
            "parsed_document_id": parsed_document_id,
            "workspace_id": workspace_id,
        }
    ]


def test_knowledge_file_chunk_list_is_scoped_and_reports_embedding_state(monkeypatch) -> None:
    workspace_id = UUID("00000000-0000-0000-0000-000000000001")
    knowledge_base_id = UUID("00000000-0000-0000-0000-000000000002")
    file_id = UUID("00000000-0000-0000-0000-000000000003")
    chunk_id = UUID("00000000-0000-0000-0000-000000000004")
    calls: list[tuple[str, dict[str, object]]] = []

    class ChunkListConnection:
        async def execute(self, query, params):
            sql = str(query)
            calls.append((sql, params))
            if "COUNT(*)" in sql:
                return FakeUploadResult(row={"chunk_count": 1})
            return FakeUploadResult(
                rows=[
                    {
                        "chunk_id": chunk_id,
                        "chunk_index": 0,
                        "content": "chair costs ten",
                        "is_embedded": False,
                        "embedding_model": None,
                    }
                ]
            )

    permission: list[object] = []

    async def fake_require_permission(*args):
        permission.extend(args[1:])

    monkeypatch.setattr(
        "app.api.routes.knowledge_files.require_knowledge_base_permission",
        fake_require_permission,
    )
    monkeypatch.setattr(
        "app.api.routes.knowledge_files.get_db_connection",
        upload_connection_context(ChunkListConnection()),
    )

    result = asyncio.run(
        list_knowledge_file_chunks(
            knowledge_base_id=knowledge_base_id,
            file_id=file_id,
            workspace_id=workspace_id,
            offset=0,
            limit=20,
            current_user=AuthenticatedUser(user_id="chunk-reader"),
            settings=Settings(knowledge_ingestion_enabled=True),
        )
    )

    assert isinstance(result, KnowledgeChunkListResponse)
    assert result.total == 1
    assert result.items[0].chunk_id == str(chunk_id)
    assert result.items[0].is_embedded is False
    assert permission == [workspace_id, knowledge_base_id, "manage"]
    assert all(params["file_id"] == file_id for _sql, params in calls)
    assert all(params["knowledge_base_id"] == knowledge_base_id for _sql, params in calls)
    assert all(params["workspace_id"] == workspace_id for _sql, params in calls)


def test_embedding_route_only_embeds_selected_chunks_for_the_requested_file(
    monkeypatch,
) -> None:
    workspace_id = UUID("00000000-0000-0000-0000-000000000001")
    knowledge_base_id = UUID("00000000-0000-0000-0000-000000000002")
    file_id = UUID("00000000-0000-0000-0000-000000000003")
    chunk_ids = [
        UUID("00000000-0000-0000-0000-000000000004"),
        UUID("00000000-0000-0000-0000-000000000005"),
    ]
    update_params: list[dict[str, object]] = []
    selected_ids: list[UUID] = []

    class ChunkEmbeddingConnection:
        async def execute(self, query, params):
            sql = str(query)
            if sql.lstrip().startswith("SELECT"):
                selected_ids.extend(params["chunk_ids"])
                return FakeUploadResult(
                    rows=[
                        {"chunk_id": chunk_ids[0], "content": "chair costs ten"},
                        {"chunk_id": chunk_ids[1], "content": "table costs twenty"},
                    ]
                )
            update_params.append(params)
            return FakeUploadResult(scalar_value=params["chunk_id"])

        def begin(self):
            return FakeTransaction()

    async def fake_embed_texts(texts, _settings):
        assert texts == ["chair costs ten", "table costs twenty"]
        return [[0.25] * 1024, [0.5] * 1024]

    permission: list[object] = []

    async def fake_require_permission(*args):
        permission.extend(args[1:])

    monkeypatch.setattr("app.api.routes.knowledge_files.embed_texts", fake_embed_texts)
    monkeypatch.setattr(
        "app.api.routes.knowledge_files.require_knowledge_base_permission",
        fake_require_permission,
    )
    monkeypatch.setattr(
        "app.api.routes.knowledge_files.get_db_connection",
        upload_connection_context(ChunkEmbeddingConnection()),
    )

    result = asyncio.run(
        embed_knowledge_file_chunks(
            knowledge_base_id=knowledge_base_id,
            file_id=file_id,
            payload=KnowledgeChunkEmbeddingRequest(chunkIds=chunk_ids),
            workspace_id=workspace_id,
            current_user=AuthenticatedUser(user_id="chunk-manager"),
            settings=Settings(
                knowledge_ingestion_enabled=True,
                knowledge_embeddings_enabled=True,
                embedding_model="qwen3.7-text-embedding",
                embedding_api_key="test-key",
            ),
        )
    )

    assert result.embedded_count == 2
    assert result.embedding_model == "qwen3.7-text-embedding"
    assert result.dimensions == 1024
    assert selected_ids == chunk_ids
    assert permission == [workspace_id, knowledge_base_id, "manage"]
    assert [params["chunk_id"] for params in update_params] == chunk_ids
    assert all(params["file_id"] == file_id for params in update_params)
    assert all(params["embedding_model"] == "qwen3.7-text-embedding" for params in update_params)


def test_embedding_route_rejects_selected_chunks_outside_the_file_without_provider_call(
    monkeypatch,
) -> None:
    workspace_id = UUID("00000000-0000-0000-0000-000000000001")
    knowledge_base_id = UUID("00000000-0000-0000-0000-000000000002")
    file_id = UUID("00000000-0000-0000-0000-000000000003")
    requested_ids = [
        UUID("00000000-0000-0000-0000-000000000004"),
        UUID("00000000-0000-0000-0000-000000000005"),
    ]
    provider_calls: list[list[str]] = []

    class ScopedChunkConnection:
        async def execute(self, _query, _params):
            return FakeUploadResult(
                rows=[{"chunk_id": requested_ids[0], "content": "in-scope chunk"}]
            )

    async def fake_embed_texts(texts, _settings):
        provider_calls.append(texts)
        return [[0.1] * 1024 for _ in texts]

    async def fake_require_permission(*_args):
        return None

    monkeypatch.setattr("app.api.routes.knowledge_files.embed_texts", fake_embed_texts)
    monkeypatch.setattr(
        "app.api.routes.knowledge_files.require_knowledge_base_permission",
        fake_require_permission,
    )
    monkeypatch.setattr(
        "app.api.routes.knowledge_files.get_db_connection",
        upload_connection_context(ScopedChunkConnection()),
    )

    with pytest.raises(Exception) as error:
        asyncio.run(
            embed_knowledge_file_chunks(
                knowledge_base_id=knowledge_base_id,
                file_id=file_id,
                payload=KnowledgeChunkEmbeddingRequest(chunkIds=requested_ids),
                workspace_id=workspace_id,
                current_user=AuthenticatedUser(user_id="chunk-manager"),
                settings=Settings(
                    knowledge_ingestion_enabled=True,
                    knowledge_embeddings_enabled=True,
                    embedding_api_key="test-key",
                ),
            )
        )

    assert getattr(error.value, "status_code", None) == 404
    assert provider_calls == []


def test_ingestion_migration_is_idempotent_and_contains_chunk_table() -> None:
    sql = MIGRATION_PATH.read_text(encoding="utf-8")

    assert 'CREATE TABLE IF NOT EXISTS "KnowledgeFile"' in sql
    assert 'CREATE TABLE IF NOT EXISTS "KnowledgeChunk"' in sql
    assert '"workspaceId" = :workspace_id' in str(FILE_LIST_QUERY)
    assert ":storage_provider" in str(FILE_INSERT_QUERY)

    parsed_migration = (
        Path(__file__).parents[1] / "migrations" / "0013_knowledge_parsed_documents.sql"
    ).read_text(encoding="utf-8")
    assert 'CREATE TABLE IF NOT EXISTS "KnowledgeParsedDocument"' in parsed_migration
    assert '"document" jsonb NOT NULL' in parsed_migration
    assert '"chunkStatus"' in parsed_migration
    assert '"fileId" = :file_id' in str(PARSED_DOCUMENT_BY_FILE_QUERY)


def test_upload_is_gated_until_ingestion_migration_is_applied(
    ingestion_disabled_settings: Settings,
) -> None:
    response = client.post(
        "/api/v1/knowledge-bases/00000000-0000-0000-0000-000000000002/files",
        params={"workspace_id": "00000000-0000-0000-0000-000000000001"},
        files={"file": ("data.csv", b"name,price\nchair,10", "text/csv")},
    )

    assert response.status_code == 409
    assert response.json()["code"] == "knowledge_ingestion:disabled"


def test_upload_route_persists_workspace_scoped_file_without_scheduling_processing(
    monkeypatch,
) -> None:
    workspace_id = UUID("00000000-0000-0000-0000-000000000001")
    knowledge_base_id = UUID("00000000-0000-0000-0000-000000000002")
    user_id = UUID("00000000-0000-0000-0000-000000000010")
    settings = Settings(
        knowledge_ingestion_enabled=True,
        knowledge_storage_dir="storage/knowledge",
    )
    connection = FakeUploadConnection()
    storage = FakeUploadStorage()
    permission: dict[str, object] = {}

    async def fake_require_permission(
        _current_user,
        requested_workspace_id,
        requested_knowledge_base_id,
        requested_permission,
    ):
        permission.update(
            {
                "workspace_id": requested_workspace_id,
                "knowledge_base_id": requested_knowledge_base_id,
                "permission": requested_permission,
            }
        )

    monkeypatch.setattr(
        "app.api.routes.knowledge_files.require_knowledge_base_permission",
        fake_require_permission,
    )
    monkeypatch.setattr(
        "app.api.routes.knowledge_files.get_knowledge_storage",
        lambda _settings: storage,
    )
    monkeypatch.setattr(
        "app.api.routes.knowledge_files.get_db_connection",
        upload_connection_context(connection),
    )

    result = asyncio.run(
        upload_knowledge_file(
            knowledge_base_id=knowledge_base_id,
            file=UploadFile(
                file=io.BytesIO(b"name,price\nchair,10"),
                filename="../data.csv",
                headers=Headers({"content-type": "text/csv"}),
            ),
            workspace_id=workspace_id,
            current_user=AuthenticatedUser(user_id=str(user_id)),
            settings=settings,
        )
    )

    assert result["file"].workspace_id == str(workspace_id)
    assert result["file"].knowledge_base_id == str(knowledge_base_id)
    assert permission == {
        "workspace_id": workspace_id,
        "knowledge_base_id": knowledge_base_id,
        "permission": "manage",
    }
    assert len(storage.uploads) == 1
    storage_key, stored_content = storage.uploads[0]
    assert storage_key.startswith(f"{workspace_id}/{knowledge_base_id}/")
    assert storage_key.endswith("-data.csv")
    assert stored_content == b"name,price\nchair,10"
    assert connection.calls[0][1]["workspace_id"] == workspace_id
    assert connection.calls[1][1]["workspace_id"] == workspace_id


def test_parse_route_claims_file_and_schedules_processing(monkeypatch) -> None:
    workspace_id = UUID("00000000-0000-0000-0000-000000000001")
    knowledge_base_id = UUID("00000000-0000-0000-0000-000000000002")
    file_id = UUID("00000000-0000-0000-0000-000000000003")
    connection = FakeUploadConnection()
    connection.file_row.update(
        {
            "file_id": file_id,
            "knowledge_base_id": knowledge_base_id,
            "workspace_id": workspace_id,
            "status": "pending",
        }
    )
    background_tasks = FakeBackgroundTasks()

    class ParseConnection(FakeUploadConnection):
        async def execute(self, query: object, params: dict[str, object]) -> FakeUploadResult:
            sql = str(query)
            if 'UPDATE "KnowledgeFile"' in sql and 'RETURNING "id"' in sql:
                self.calls.append((sql, params))
                self.file_row["status"] = "processing"
                return FakeUploadResult(scalar_value=file_id)
            return await super().execute(query, params)

    parse_connection = ParseConnection()
    parse_connection.file_row = connection.file_row

    async def fake_require_permission(*_args, **_kwargs):
        return SimpleNamespace(role="owner")

    monkeypatch.setattr(
        "app.api.routes.knowledge_files.require_knowledge_base_permission",
        fake_require_permission,
    )
    monkeypatch.setattr(
        "app.api.routes.knowledge_files.get_db_connection",
        upload_connection_context(parse_connection),
    )

    result = asyncio.run(
        parse_knowledge_file(
            knowledge_base_id=knowledge_base_id,
            file_id=file_id,
            background_tasks=background_tasks,
            workspace_id=workspace_id,
            current_user=AuthenticatedUser(user_id=str(UUID("00000000-0000-0000-0000-000000000010"))),
            settings=Settings(knowledge_ingestion_enabled=True),
        )
    )

    assert result["file"].status == "processing"
    assert background_tasks.tasks == [(process_knowledge_file, (file_id, workspace_id))]


def test_parsed_document_is_readable_and_chunking_is_a_separate_manual_action(
    monkeypatch,
) -> None:
    workspace_id = UUID("00000000-0000-0000-0000-000000000001")
    knowledge_base_id = UUID("00000000-0000-0000-0000-000000000002")
    file_id = UUID("00000000-0000-0000-0000-000000000003")
    parsed_document_id = UUID("00000000-0000-0000-0000-000000000004")
    document = parse_document(
        "data.csv",
        b"name,price\nchair,10",
        file_id=str(file_id),
        file_hash="hash-123",
        mime_type="text/csv",
    )

    class ParsedDocumentConnection(FakeUploadConnection):
        def __init__(self) -> None:
            super().__init__()
            self.parsed_row = {
                "block_count": len(document.blocks),
                "chunk_error_message": None,
                "chunk_status": "pending",
                "created_at": datetime(2026, 8, 17, 12, 30),
                "document": document.model_dump(by_alias=True),
                "file_hash": "hash-123",
                "file_id": file_id,
                "parsed_document_id": parsed_document_id,
                "knowledge_base_id": knowledge_base_id,
                "parser": document.parser,
                "parser_version": document.parser_version,
                "schema_version": document.schema_version,
                "updated_at": datetime(2026, 8, 17, 12, 30),
                "warning_count": len(document.warnings),
                "workspace_id": workspace_id,
            }

        async def execute(self, query: object, params: dict[str, object]) -> FakeUploadResult:
            sql = str(query)
            self.calls.append((sql, params))
            if 'UPDATE "KnowledgeParsedDocument"' in sql:
                self.parsed_row["chunk_status"] = "processing"
                return FakeUploadResult(scalar_value=parsed_document_id)
            if 'SELECT' in sql and 'KnowledgeParsedDocument' in sql:
                return FakeUploadResult(row=self.parsed_row)
            if 'COUNT(*)' in sql and 'KnowledgeChunk' in sql:
                return FakeUploadResult(row={"chunk_count": 0})
            return await super().execute(query, params)

    connection = ParsedDocumentConnection()
    background_tasks = FakeBackgroundTasks()

    async def fake_require_permission(*_args, **_kwargs):
        return SimpleNamespace(role="owner")

    monkeypatch.setattr(
        "app.api.routes.knowledge_files.require_knowledge_base_permission",
        fake_require_permission,
    )
    monkeypatch.setattr(
        "app.api.routes.knowledge_files.get_db_connection",
        upload_connection_context(connection),
    )
    settings = Settings(knowledge_ingestion_enabled=True)
    user = AuthenticatedUser(user_id=str(UUID("00000000-0000-0000-0000-000000000010")))

    parsed_result = asyncio.run(
        get_parsed_knowledge_document(
            knowledge_base_id=knowledge_base_id,
            file_id=file_id,
            workspace_id=workspace_id,
            current_user=user,
            settings=settings,
        )
    )
    assert isinstance(parsed_result, KnowledgeParsedDocumentResponse)
    assert parsed_result.parsed_document.blocks
    assert parsed_result.chunk_status == "pending"

    chunk_result = asyncio.run(
        materialize_knowledge_file_chunks(
            knowledge_base_id=knowledge_base_id,
            file_id=file_id,
            background_tasks=background_tasks,
            workspace_id=workspace_id,
            current_user=user,
            settings=settings,
        )
    )
    assert isinstance(chunk_result, KnowledgeParsedDocumentResponse)
    assert chunk_result.chunk_status == "processing"
    assert background_tasks.tasks == [
        (materialize_knowledge_chunks, (file_id, workspace_id, parsed_document_id))
    ]


def test_knowledge_base_parsed_document_list_returns_all_parsed_files(
    monkeypatch,
) -> None:
    workspace_id = UUID("00000000-0000-0000-0000-000000000001")
    knowledge_base_id = UUID("00000000-0000-0000-0000-000000000002")
    first_file_id = UUID("00000000-0000-0000-0000-000000000003")
    second_file_id = UUID("00000000-0000-0000-0000-000000000004")
    first_document = parse_document(
        "first.csv",
        b"name,price\nchair,10",
        file_id=str(first_file_id),
        file_hash="hash-first",
        mime_type="text/csv",
    )
    second_document = parse_document(
        "second.txt",
        b"A product description.",
        file_id=str(second_file_id),
        file_hash="hash-second",
        mime_type="text/plain",
    )
    rows = [
        {
            "file_byte_size": 20,
            "chunk_count": 2,
            "chunk_error_message": None,
            "chunk_status": "ready",
            "created_at": datetime(2026, 8, 17, 12, 30),
            "document": first_document.model_dump(by_alias=True),
            "file_hash": "hash-first",
            "file_id": first_file_id,
            "file_mime_type": "text/csv",
            "file_name": "first.csv",
            "file_status": "ready",
            "parsed_document_id": UUID("00000000-0000-0000-0000-000000000005"),
            "updated_at": datetime(2026, 8, 17, 12, 30),
        },
        {
            "file_byte_size": 21,
            "chunk_count": 0,
            "chunk_error_message": None,
            "chunk_status": "pending",
            "created_at": datetime(2026, 8, 17, 12, 31),
            "document": second_document.model_dump(by_alias=True),
            "file_hash": "hash-second",
            "file_id": second_file_id,
            "file_mime_type": "text/plain",
            "file_name": "second.txt",
            "file_status": "ready",
            "parsed_document_id": UUID("00000000-0000-0000-0000-000000000006"),
            "updated_at": datetime(2026, 8, 17, 12, 31),
        },
    ]

    class ParsedDocumentListConnection:
        async def execute(self, query: object, params: dict[str, object]):
            sql = str(query)
            if sql.lstrip().startswith("SELECT COUNT(*)"):
                return FakeUploadResult(row={"document_count": len(rows)})
            return FakeUploadResult(rows=rows)

    async def fake_require_permission(*_args, **_kwargs):
        return SimpleNamespace(role="owner")

    monkeypatch.setattr(
        "app.api.routes.knowledge_files.require_knowledge_base_permission",
        fake_require_permission,
    )
    monkeypatch.setattr(
        "app.api.routes.knowledge_files.get_db_connection",
        upload_connection_context(ParsedDocumentListConnection()),
    )

    result = asyncio.run(
        list_parsed_knowledge_documents(
            knowledge_base_id=knowledge_base_id,
            workspace_id=workspace_id,
            current_user=AuthenticatedUser(
                user_id="00000000-0000-0000-0000-000000000010"
            ),
            settings=Settings(knowledge_ingestion_enabled=True),
        )
    )

    assert isinstance(result, KnowledgeParsedDocumentListResponse)
    assert result.total == 2
    assert [item.file_id for item in result.items] == [
        str(first_file_id),
        str(second_file_id),
    ]
    assert result.items[0].file_name == "first.csv"
    assert result.items[0].parsed_document.blocks
    assert result.items[1].chunk_status == "pending"


def test_duplicate_upload_removes_object_after_database_conflict(monkeypatch) -> None:
    workspace_id = UUID("00000000-0000-0000-0000-000000000001")
    knowledge_base_id = UUID("00000000-0000-0000-0000-000000000002")
    settings = Settings(knowledge_ingestion_enabled=True)
    connection = FakeUploadConnection(inserted=False)
    storage = FakeUploadStorage()

    async def fake_require_permission(*_args, **_kwargs):
        return SimpleNamespace(role="owner")

    monkeypatch.setattr(
        "app.api.routes.knowledge_files.require_knowledge_base_permission",
        fake_require_permission,
    )
    monkeypatch.setattr(
        "app.api.routes.knowledge_files.get_knowledge_storage",
        lambda _settings: storage,
    )
    monkeypatch.setattr(
        "app.api.routes.knowledge_files.get_db_connection",
        upload_connection_context(connection),
    )

    with pytest.raises(Exception) as error:
        asyncio.run(
            upload_knowledge_file(
                knowledge_base_id=knowledge_base_id,
                file=UploadFile(
                    file=io.BytesIO(b"name,price\nchair,10"),
                    filename="data.csv",
                    headers=Headers({"content-type": "text/csv"}),
                ),
                workspace_id=workspace_id,
                current_user=AuthenticatedUser(
                    user_id="00000000-0000-0000-0000-000000000010"
                ),
                settings=settings,
            )
        )

    assert getattr(error.value, "status_code", None) == 409
    assert len(storage.uploads) == 1
    assert storage.deleted == [storage.uploads[0][0]]
