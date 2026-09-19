"""Deterministic document parsing for knowledge-file ingestion and agent tools.

Parsers in this module only extract evidence from a source file.  They do not
infer business entities, execute formulas, or write to a database.  A later
agent/import stage can use the stable ``ParsedDocument`` output to propose
business records and vector chunks.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from functools import lru_cache
from pathlib import Path
from threading import Lock
from typing import Any, Literal, Protocol

from openpyxl import load_workbook
from pptx import Presentation
from pydantic import BaseModel, ConfigDict, Field
from rapidocr_pdf import RapidOCRPDF

SUPPORTED_EXTENSIONS = {
    ".csv": "text/csv",
    ".json": "application/json",
    ".md": "text/markdown",
    ".pdf": "application/pdf",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".txt": "text/plain",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

FILE_SIGNATURES = {
    ".pdf": (b"%PDF-",),
    ".pptx": (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08"),
    ".xlsx": (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08"),
}

PARSED_DOCUMENT_SCHEMA_VERSION = "parsed-document.v1"
MAX_PARSE_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 10_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 512 * 1024 * 1024
_PDF_OCR_LOCK = Lock()

ContentType = Literal["text", "structured", "mixed"]
OutputFormat = Literal["text", "structured"]
LocatorValue = str | int


class ParsedBlock(BaseModel):
    """One deterministic evidence block extracted from a source file."""

    model_config = ConfigDict(extra="forbid", populate_by_name=True)

    block_id: str = Field(alias="blockId")
    kind: str
    text: str
    data: Any | None = None
    locator: dict[str, LocatorValue]
    extraction_method: str = Field(default="native", alias="extractionMethod")


class ParsedDocument(BaseModel):
    """Stable intermediate representation shared by ingestion and agent tools."""

    model_config = ConfigDict(extra="forbid", populate_by_name=True)

    schema_version: str = Field(
        default=PARSED_DOCUMENT_SCHEMA_VERSION,
        alias="schemaVersion",
    )
    file_id: str | None = Field(default=None, alias="fileId")
    file_hash: str = Field(alias="fileHash")
    original_name: str = Field(alias="originalName")
    mime_type: str = Field(alias="mimeType")
    parser: str
    parser_version: str = Field(alias="parserVersion")
    content_type: ContentType = Field(alias="contentType")
    blocks: list[ParsedBlock]
    warnings: list[str] = Field(default_factory=list)
    truncated: bool = False
    total_blocks: int | None = Field(default=None, alias="totalBlocks")
    next_cursor: int | None = Field(default=None, alias="nextCursor")


@dataclass(frozen=True)
class _RawBlock:
    kind: str
    text: str
    locator: dict[str, LocatorValue]
    data: Any | None = None
    extraction_method: str = "native"


@dataclass(frozen=True)
class _ParserOutput:
    parser: str
    parser_version: str
    content_type: ContentType
    blocks: list[_RawBlock]
    warnings: list[str]


class _Parser(Protocol):
    def __call__(self, content: bytes) -> _ParserOutput: ...


def _decode_text(content: bytes) -> str:
    return content.decode("utf-8-sig", errors="replace").replace("\r\n", "\n").replace(
        "\r", "\n"
    )


def _paragraphs(text: str) -> list[tuple[str, int, int]]:
    """Return non-empty paragraph text and its one-based line range."""
    lines = text.split("\n")
    result: list[tuple[str, int, int]] = []
    start: int | None = None
    values: list[str] = []

    def flush(end: int) -> None:
        nonlocal start, values
        if start is not None and values:
            result.append(("\n".join(values).strip(), start, end))
        start = None
        values = []

    for index, line in enumerate(lines, start=1):
        if line.strip():
            if start is None:
                start = index
            values.append(line.rstrip())
        elif start is not None:
            flush(index - 1)
    if start is not None:
        flush(len(lines))
    return [(value, start, end) for value, start, end in result if value]


def _parse_text(content: bytes) -> _ParserOutput:
    text = _decode_text(content)
    stripped = text.strip()
    blocks = (
        [
            _RawBlock(
                kind="paragraph",
                text=stripped,
                locator={"lineStart": 1, "lineEnd": len(text.splitlines())},
            )
        ]
        if stripped
        else []
    )
    warnings = [] if blocks else ["The text file contains no extractable text."]
    return _ParserOutput("text", "text-1", "text", blocks, warnings)


_MARKDOWN_HEADING = re.compile(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$")
_MARKDOWN_TABLE_ROW = re.compile(r"^\s*\|?(.*?)\|\s*$")
_MARKDOWN_SEPARATOR_CELL = re.compile(r"^\s*:?-{1,}:?\s*$")


def _markdown_cells(line: str) -> list[str]:
    match = _MARKDOWN_TABLE_ROW.match(line)
    if not match:
        return []
    return [cell.strip() for cell in match.group(1).split("|")]


def _parse_markdown(content: bytes) -> _ParserOutput:
    text = _decode_text(content)
    lines = text.split("\n")
    blocks: list[_RawBlock] = []
    warnings: list[str] = []
    index = 0
    has_table = False

    while index < len(lines):
        line = lines[index]
        heading = _MARKDOWN_HEADING.match(line)
        if heading:
            blocks.append(
                _RawBlock(
                    kind="heading",
                    text=heading.group(2).strip(),
                    locator={"lineStart": index + 1, "lineEnd": index + 1},
                    data={"level": len(heading.group(1))},
                )
            )
            index += 1
            continue

        if (
            index + 1 < len(lines)
            and "|" in line
            and all(
                _MARKDOWN_SEPARATOR_CELL.fullmatch(cell)
                for cell in _markdown_cells(lines[index + 1])
            )
        ):
            headers = _markdown_cells(line)
            rows: list[dict[str, str]] = []
            row_text: list[str] = ["\t".join(headers)]
            end = index + 1
            cursor = index + 2
            while cursor < len(lines) and lines[cursor].strip() and "|" in lines[cursor]:
                values = _markdown_cells(lines[cursor])
                row = _row_mapping(headers, values)
                rows.append(row)
                row_text.append("\t".join(values))
                end = cursor
                cursor += 1
            blocks.append(
                _RawBlock(
                    kind="table",
                    text="\n".join(row_text),
                    locator={"lineStart": index + 1, "lineEnd": end + 1},
                    data={"headers": headers, "rows": rows},
                )
            )
            has_table = True
            index = cursor
            continue

        if not line.strip():
            index += 1
            continue

        start = index
        paragraph_lines = [line.rstrip()]
        index += 1
        while index < len(lines) and lines[index].strip():
            if _MARKDOWN_HEADING.match(lines[index]):
                break
            if index + 1 < len(lines) and "|" in lines[index] and "|" in lines[index + 1]:
                break
            paragraph_lines.append(lines[index].rstrip())
            index += 1
        blocks.append(
            _RawBlock(
                kind="paragraph",
                text="\n".join(paragraph_lines).strip(),
                locator={"lineStart": start + 1, "lineEnd": index},
            )
        )

    if not blocks:
        warnings.append("The Markdown file contains no extractable text.")
    return _ParserOutput(
        "markdown",
        "markdown-1",
        "mixed" if has_table else "text",
        blocks,
        warnings,
    )


def _row_mapping(headers: list[str], values: list[Any]) -> dict[str, Any]:
    mapping: dict[str, Any] = {}
    for index, header in enumerate(headers):
        mapping[header] = values[index] if index < len(values) else ""
    for index, value in enumerate(values[len(headers) :], start=len(headers) + 1):
        mapping[f"column_{index}"] = value
    return mapping


def _unique_headers(values: list[Any]) -> list[str]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = str(value).strip() if value is not None else ""
        base = base or f"column_{index}"
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return headers


def _parse_csv(content: bytes) -> _ParserOutput:
    reader = csv.reader(io.StringIO(_decode_text(content)))
    rows = list(reader)
    nonempty = [
        (index + 1, row)
        for index, row in enumerate(rows)
        if any(cell.strip() for cell in row)
    ]
    if not nonempty:
        return _ParserOutput(
            "csv",
            "csv-1",
            "structured",
            [],
            ["The CSV file contains no rows."],
        )

    header_line, header_values = nonempty[0]
    headers = _unique_headers(header_values)
    blocks = [
        _RawBlock(
            kind="table_header",
            text="\t".join(headers),
            locator={"row": header_line},
            data={"headers": headers},
        )
    ]
    for row_number, values in nonempty[1:]:
        blocks.append(
            _RawBlock(
                kind="table_row",
                text="\t".join(value.strip() for value in values),
                locator={"row": row_number},
                data=_row_mapping(headers, values),
            )
        )
    return _ParserOutput("csv", "csv-1", "structured", blocks, [])


def _validate_zip_payload(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise ValueError("The document archive contains too many entries.")
            uncompressed_size = sum(member.file_size for member in members)
            if uncompressed_size > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise ValueError("The document archive expands beyond the safe limit.")
    except zipfile.BadZipFile as error:
        raise ValueError("The uploaded document is not a valid Office archive.") from error


def _safe_cell_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _parse_xlsx(content: bytes) -> _ParserOutput:
    _validate_zip_payload(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError("The XLSX document could not be parsed.") from error

    blocks: list[_RawBlock] = []
    warnings: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            sheet_name = worksheet.title
            blocks.append(
                _RawBlock(
                    kind="sheet",
                    text=f"[Sheet: {sheet_name}]",
                    locator={"sheet": sheet_name},
                    data={"sheet": sheet_name},
                )
            )
            rows = [
                (row_number, [_safe_cell_value(value) for value in values])
                for row_number, values in enumerate(worksheet.iter_rows(values_only=True), start=1)
                if any(value is not None and str(value).strip() for value in values)
            ]
            if not rows:
                warnings.append(f"Sheet '{sheet_name}' contains no extractable rows.")
                continue
            header_number, header_values = rows[0]
            headers = _unique_headers(header_values)
            blocks.append(
                _RawBlock(
                    kind="table_header",
                    text="\t".join(headers),
                    locator={"sheet": sheet_name, "row": header_number},
                    data={"headers": headers},
                )
            )
            for row_number, values in rows[1:]:
                blocks.append(
                    _RawBlock(
                        kind="table_row",
                        text="\t".join(
                            "" if value is None else str(value).strip() for value in values
                        ),
                        locator={"sheet": sheet_name, "row": row_number},
                        data=_row_mapping(headers, values),
                    )
                )
    finally:
        workbook.close()
    return _ParserOutput("xlsx", "xlsx-1", "structured", blocks, warnings)


@lru_cache(maxsize=1)
def _get_pdf_extractor() -> RapidOCRPDF:
    """Create one bounded RapidOCRPDF instance and reuse its loaded models."""
    return RapidOCRPDF(
        ocr_params={
            "EngineConfig.onnxruntime.intra_op_num_threads": 2,
            "EngineConfig.onnxruntime.inter_op_num_threads": 1,
        }
    )


def _parse_pdf_page_result(result: object) -> tuple[int, str, float | None]:
    if not isinstance(result, (list, tuple)) or len(result) < 2:
        raise ValueError("The PDF OCR returned an invalid page result.")
    try:
        page_number = int(result[0]) + 1
    except (TypeError, ValueError) as error:
        raise ValueError("The PDF OCR returned an invalid page number.") from error

    page_text = "" if result[1] is None else str(result[1])
    confidence: float | None = None
    raw_confidence = result[2] if len(result) >= 3 else None
    if raw_confidence not in (None, "", "N/A", "n/a"):
        try:
            confidence = float(raw_confidence)
        except (TypeError, ValueError) as error:
            raise ValueError("The PDF OCR returned an invalid confidence score.") from error
    return page_number, page_text, confidence


def _parse_pdf(content: bytes) -> _ParserOutput:
    try:
        # The wrapper uses native PDF text extraction where available and
        # falls back to RapidOCR for scanned pages.  Serializing calls keeps
        # the two-vCPU deployment from loading competing OCR jobs at once.
        with _PDF_OCR_LOCK:
            page_results = _get_pdf_extractor()(content, force_ocr=False)
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("The PDF document could not be parsed with RapidOCR.") from error

    blocks: list[_RawBlock] = []
    warnings: list[str] = []
    try:
        results = list(page_results or [])
    except TypeError as error:
        raise ValueError("The PDF OCR returned an invalid result.") from error

    for result in results:
        page_number, page_text, confidence = _parse_pdf_page_result(result)
        paragraphs = _paragraphs(page_text.replace("\r\n", "\n").replace("\r", "\n"))
        if not paragraphs:
            warnings.append(
                f"Page {page_number} has no extractable text after PDF text extraction and OCR."
            )
        for paragraph_number, (value, _start, _end) in enumerate(paragraphs, start=1):
            blocks.append(
                _RawBlock(
                    kind="paragraph",
                    text=value,
                    locator={"page": page_number, "paragraph": paragraph_number},
                    data={"confidence": confidence} if confidence is not None else None,
                    extraction_method="rapidocr_pdf",
                )
            )
    if not blocks and not warnings:
        warnings.append("The PDF contains no extractable text after PDF text extraction and OCR.")
    return _ParserOutput("pdf-rapidocr", "pdf-rapidocr-1", "text", blocks, warnings)


def _parse_pptx_shape(shape: object, slide_number: int, shape_number: int) -> list[_RawBlock]:
    locator_prefix: dict[str, LocatorValue] = {"slide": slide_number, "shape": shape_number}
    if getattr(shape, "has_table", False):
        table = shape.table
        blocks: list[_RawBlock] = []
        for row_number, row in enumerate(table.rows, start=1):
            values = [cell.text.strip() for cell in row.cells]
            if not any(values):
                continue
            locator = {**locator_prefix, "row": row_number}
            if row_number == 1:
                blocks.append(
                    _RawBlock(
                        kind="table_header",
                        text="\t".join(values),
                        locator=locator,
                        data={"headers": values},
                    )
                )
            else:
                headers = [cell.text.strip() for cell in table.rows[0].cells]
                blocks.append(
                    _RawBlock(
                        kind="table_row",
                        text="\t".join(values),
                        locator=locator,
                        data=_row_mapping(headers, values),
                    )
                )
        return blocks
    if getattr(shape, "has_text_frame", False):
        text = shape.text.strip()
        return [
            _RawBlock(kind="text_box", text=text, locator=locator_prefix)
        ] if text else []
    nested_shapes = getattr(shape, "shapes", None)
    if nested_shapes is None:
        return []
    blocks = []
    for nested_number, nested_shape in enumerate(nested_shapes, start=1):
        for block in _parse_pptx_shape(nested_shape, slide_number, shape_number):
            blocks.append(
                _RawBlock(
                    kind=block.kind,
                    text=block.text,
                    locator={**block.locator, "nestedShape": nested_number},
                    data=block.data,
                    extraction_method=block.extraction_method,
                )
            )
    return blocks


def _parse_pptx(content: bytes) -> _ParserOutput:
    _validate_zip_payload(content)
    try:
        presentation = Presentation(io.BytesIO(content))
    except Exception as error:
        raise ValueError("The PPTX document could not be parsed.") from error

    blocks: list[_RawBlock] = []
    warnings: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_blocks: list[_RawBlock] = []
        for shape_number, shape in enumerate(slide.shapes, start=1):
            slide_blocks.extend(_parse_pptx_shape(shape, slide_number, shape_number))
        if slide_blocks:
            blocks.append(
                _RawBlock(
                    kind="slide",
                    text="",
                    locator={"slide": slide_number},
                    data={"slide": slide_number},
                )
            )
            blocks.extend(slide_blocks)
        else:
            warnings.append(f"Slide {slide_number} contains no extractable text.")
    return _ParserOutput("pptx", "pptx-1", "mixed", blocks, warnings)


_PARSERS: dict[str, _Parser] = {
    ".csv": _parse_csv,
    ".json": lambda content: _parse_json(content),
    ".md": _parse_markdown,
    ".pdf": _parse_pdf,
    ".pptx": _parse_pptx,
    ".txt": _parse_text,
    ".xlsx": _parse_xlsx,
}


def _parse_json(content: bytes) -> _ParserOutput:
    try:
        value = json.loads(_decode_text(content))
    except (UnicodeDecodeError, json.JSONDecodeError) as error:
        raise ValueError("The JSON document could not be parsed.") from error

    values = list(enumerate(value)) if isinstance(value, list) else [(None, value)]
    blocks = []
    for index, item in values:
        path = f"$[{index}]" if index is not None else "$"
        blocks.append(
            _RawBlock(
                kind="json_value" if not isinstance(item, (dict, list)) else "json_record",
                text=json.dumps(item, ensure_ascii=False, indent=2, sort_keys=True),
                locator={"jsonPath": path},
                data=item,
            )
        )
    warnings = [] if blocks else ["The JSON array contains no records."]
    return _ParserOutput("json", "json-1", "structured", blocks, warnings)


def _stable_block_id(file_hash: str, parser_version: str, index: int, block: _RawBlock) -> str:
    payload = json.dumps(
        {
            "fileHash": file_hash,
            "parserVersion": parser_version,
            "index": index,
            "kind": block.kind,
            "text": block.text,
            "locator": block.locator,
        },
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return f"block-{hashlib.sha256(payload.encode('utf-8')).hexdigest()[:24]}"


def parse_document(
    filename: str,
    content: bytes,
    *,
    file_id: str | None = None,
    file_hash: str | None = None,
    mime_type: str | None = None,
) -> ParsedDocument:
    """Parse one supported source into the versioned intermediate format."""
    if len(content) > MAX_PARSE_BYTES:
        raise ValueError("The document is larger than the parser safety limit.")
    extension = Path(filename).suffix.lower()
    parser = _PARSERS.get(extension)
    if parser is None:
        raise ValueError("Unsupported knowledge file type.")

    source_hash = file_hash or hashlib.sha256(content).hexdigest()
    try:
        output = parser(content)
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("The document could not be parsed.") from error

    blocks = [
        ParsedBlock(
            blockId=_stable_block_id(source_hash, output.parser_version, index, block),
            kind=block.kind,
            text=block.text,
            data=block.data,
            locator=block.locator,
            extractionMethod=block.extraction_method,
        )
        for index, block in enumerate(output.blocks)
    ]
    return ParsedDocument(
        fileId=file_id,
        fileHash=source_hash,
        originalName=filename,
        mimeType=mime_type or SUPPORTED_EXTENSIONS[extension],
        parser=output.parser,
        parserVersion=output.parser_version,
        contentType=output.content_type,
        blocks=blocks,
        warnings=output.warnings,
    )


def paginate_parsed_document(
    document: ParsedDocument,
    *,
    output: OutputFormat = "structured",
    page: int | None = None,
    sheet: str | None = None,
    slide: int | None = None,
    offset: int = 0,
    limit: int = 20,
) -> ParsedDocument:
    """Filter and page a parsed document without changing evidence ordering."""
    if offset < 0 or limit < 0:
        raise ValueError("Document pagination arguments are invalid.")

    def matches(block: ParsedBlock) -> bool:
        locator = block.locator
        return (
            (page is None or locator.get("page") == page)
            and (sheet is None or locator.get("sheet") == sheet)
            and (slide is None or locator.get("slide") == slide)
        )

    filtered = [block for block in document.blocks if matches(block)]
    selected = filtered[offset : offset + limit] if limit else []
    if output == "text":
        selected = [block.model_copy(update={"data": None}) for block in selected]
    next_cursor = (
        offset + limit
        if limit and offset + limit < len(filtered)
        else None
    )
    return document.model_copy(
        update={
            "blocks": selected,
            "total_blocks": len(filtered),
            "next_cursor": next_cursor,
            "truncated": next_cursor is not None,
        }
    )


def render_document_text(document: ParsedDocument, *, include_markers: bool = False) -> str:
    """Render blocks as text, retaining legacy sheet/slide markers when requested."""
    lines: list[str] = []
    last_slide: int | None = None
    last_shape: int | None = None
    for block in document.blocks:
        locator = block.locator
        if include_markers and "slide" in locator:
            slide_number = int(locator["slide"])
            if slide_number != last_slide:
                lines.append(f"[Slide: {slide_number}]")
                last_slide = slide_number
                last_shape = None
            if "shape" in locator:
                shape_number = int(locator["shape"])
                if shape_number != last_shape:
                    lines.append(f"[Shape: {shape_number}]")
                    last_shape = shape_number
        if block.text:
            lines.append(block.text)
    return "\n".join(line for line in lines if line)
